#!/usr/bin/env python3
"""Read CSV/XLSX keyword rows without changing data or inventing metrics."""
import argparse
import csv
import json
from pathlib import Path


def records(rows, query=(), limit=100):
    iterator = iter(rows)
    try:
        headers = next(iterator)
    except StopIteration:
        return []
    names, used = [], set()
    for index, value in enumerate(headers, start=1):
        base = str(value).strip() if value is not None and str(value).strip() else f"column_{index}"
        name = base
        suffix = 2
        while name in used:
            name = f"{base}_{suffix}"
            suffix += 1
        used.add(name)
        names.append(name)
    terms = [term.casefold() for term in query if term.strip()]
    result = []
    for number, row in enumerate(iterator, start=2):
        if not any(v is not None and str(v).strip() for v in row):
            continue
        haystack = " ".join(str(v) for v in row if v is not None).casefold()
        if terms and not any(term in haystack for term in terms):
            continue
        fields = {name: (row[i] if i < len(row) and row[i] != "" else None)
                  for i, name in enumerate(names)}
        result.append({"row": number, "fields": fields})
        if len(result) >= limit:
            break
    return result


def read_inventory(path, sheet=None, query=(), limit=100):
    if limit < 1:
        raise ValueError("Limit must be positive")
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as handle:
            result = records(csv.reader(handle), query, limit)
        sheets, chosen = ["CSV"], "CSV"
    elif path.suffix.lower() == ".xlsx":
        try:
            import openpyxl
        except ImportError as exc:
            raise ValueError("XLSX reading needs openpyxl. Use an existing spreadsheet runtime, install it in an isolated environment, or supply a CSV export.") from exc
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            sheets = workbook.sheetnames
            chosen = sheet or sheets[0]
            if chosen not in sheets:
                raise ValueError(f"Unknown sheet {chosen!r}; available: {', '.join(sheets)}")
            result = records(workbook[chosen].iter_rows(values_only=True), query, limit)
        finally:
            workbook.close()
    else:
        raise ValueError("Supported files: .csv and .xlsx")
    return {"source_file": path.name, "available_sheets": sheets, "sheet": chosen,
            "matches": result, "limit": limit,
            "notes": "First row is treated as headers. Query terms use OR substring matching across cells. Blank metrics stay null; formula cells require cached values. Values are not live-validated search data."}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("file", type=Path)
    parser.add_argument("--sheet")
    parser.add_argument("--query", action="append", default=[])
    parser.add_argument("--limit", type=int, default=100)
    args = parser.parse_args()
    try:
        result = read_inventory(args.file, args.sheet, args.query, args.limit)
    except (OSError, ValueError, UnicodeError) as exc:
        parser.exit(2, f"Cannot read inventory: {exc}\n")
    print(json.dumps(result, ensure_ascii=False, indent=2, default=str))


if __name__ == "__main__":
    main()
